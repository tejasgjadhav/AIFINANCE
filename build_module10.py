#!/usr/bin/env python3
"""Build Module 10 deck: GCC Capstone, one integrated financial AI system.
Ten simple slides; the worked example is the daily pre-market dashboard, shown by design only."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

NAVY = RGBColor(0x0F, 0x21, 0x45)
NAVY2 = RGBColor(0x0B, 0x19, 0x36)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLDD = RGBColor(0x8F, 0x71, 0x13)
BLUE = RGBColor(0x3A, 0x5C, 0x9E)
INK = RGBColor(0x1B, 0x2A, 0x41)
BG = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTE = RGBColor(0x5D, 0x6B, 0x7E)
MUTEN = RGBColor(0x8F, 0xA1, 0xC4)  # muted on navy
SUBN = RGBColor(0xC9, 0xD6, 0xEE)   # subtitle on navy
TINT1 = RGBColor(0xE9, 0xEF, 0xF8)
TINT2 = RGBColor(0xD5, 0xE1, 0xF2)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

FOOT_L = "ISBMS · PGDM 2025–27 · AGENTIC AI & ADVANCED ANALYTICS IN FINANCE"


def slide_bg(s, color):
    el = s.background.fill
    el.solid()
    el.fore_color.rgb = color


def box(s, x, y, w, h, fill=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.line.fill.background()
    sh.shadow.inherit = False
    if fill is not None:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    return sh


def txt(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, wrap=True):
    """paras: list of paragraphs; each = (align_or_None, space_before_pt, [runs])
    run = (text, font_name_or_None, size_pt, bold, color)"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in (tf.margin_left, ):
        pass
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for p_align, sp, runs in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = p_align if p_align is not None else align
        if sp:
            p.space_before = Pt(sp)
        for text, fname, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            if fname:
                r.font.name = fname
    return tb


def eyebrow(s, label):
    box(s, 0.6, 0.44, 0.13, 0.13, GOLD)
    txt(s, 0.86, 0.34, 11.83, 0.3, [(None, 0, [(label, None, 11, True, GOLD)])])


def title(s, text):
    txt(s, 0.6, 0.68, 12.13, 0.8, [(None, 0, [(text, "Georgia", 29, True, NAVY)])])


def footer(s, num, dark=False):
    num = len(prs.slides._sldIdLst)  # number slides by position, so inserts never break the sequence
    c = MUTEN if dark else MUTE
    txt(s, 0.6, 7.08, 8.0, 0.25, [(None, 0, [(FOOT_L, None, 8, False, c)])])
    txt(s, 10.13, 7.08, 2.6, 0.25, [(PP_ALIGN.RIGHT, 0, [(f"MODULE 10  ·  {num:02d}", None, 8, False, c)])],
        align=PP_ALIGN.RIGHT)


def goal_band(s, lead, rest):
    box(s, 0.6, 6.3, 12.13, 0.66, NAVY)
    box(s, 0.6, 6.3, 0.07, 0.66, GOLD)
    txt(s, 0.9, 6.3, 11.53, 0.66,
        [(None, 0, [(lead + "  ", None, 12, True, GOLD), (rest, None, 12, False, WHITE)])],
        anchor=MSO_ANCHOR.MIDDLE)


def bullet(s, x, y, w, text, size=11):
    txt(s, x, y, w, 0.55, [(None, 0, [("▪  ", None, size, True, GOLD), (text, None, size, False, INK)])])


def new_slide():
    return prs.slides.add_slide(BLANK)


def minibox(s, x, y, label, w=0.62, h=0.34, fill=NAVY, tc=WHITE):
    box(s, x, y, w, h, fill)
    txt(s, x, y, w, h, [(PP_ALIGN.CENTER, 0, [(label, None, 8, True, tc)])],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def arrow(s, x, y, ch="→", w=0.3):
    txt(s, x, y, w, 0.34, [(PP_ALIGN.CENTER, 0, [(ch, None, 13, True, GOLD)])],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def three_cards(s, cards, top=1.72, h=4.2):
    """cards: list of (fill, kick, head, body, example, useline)"""
    xs = [0.6, 4.72, 8.84]
    for x, (fill, kick, head, body, ex, use) in zip(xs, cards):
        dark = fill == NAVY
        box(s, x, top, 3.95, h, fill)
        txt(s, x + 0.26, top + 0.2, 3.4, 0.3, [(None, 0, [(kick, None, 10, True, GOLD if dark else BLUE)])])
        txt(s, x + 0.26, top + 0.52, 3.4, 0.4, [(None, 0, [(head, "Georgia", 15, True, WHITE if dark else NAVY)])])
        txt(s, x + 0.26, top + 1.05, 3.4, 1.55, [(None, 0, [(body, None, 10, False, SUBN if dark else INK)])])
        box(s, x + 0.26, top + 2.62, 3.43, 1.0, WHITE)
        txt(s, x + 0.44, top + 2.72, 3.1, 0.85, [
            (None, 0, [("EXAMPLE", None, 8, True, BLUE)]),
            (None, 2, [(ex, None, 9.5, False, INK)]),
        ])
        txt(s, x + 0.26, top + 3.7, 3.4, 0.45, [(None, 0, [(use, None, 10, True, SUBN if dark else MUTE)])])

# ---------------------------------------------------------------- SIMPLE LAYOUT HELPER
def cards(s, items, cols, top=1.6, h=None, gap=0.17, dark_last=False):
    """items: (kicker, head, body). Equal cards in a grid between top and 6.15."""
    rows = (len(items) + cols - 1) // cols
    w = (12.13 - gap * (cols - 1)) / cols
    h = h or (6.15 - top - gap * (rows - 1)) / rows
    for i, (k, hd, b) in enumerate(items):
        x = 0.6 + (i % cols) * (w + gap)
        y = top + (i // cols) * (h + gap)
        dark = dark_last and i == len(items) - 1
        box(s, x, y, w, h, NAVY if dark else WHITE)
        box(s, x, y, w, 0.05, GOLD)
        paras = []
        if k:
            paras.append((None, 0, [(k, None, 10, True, GOLD if dark else BLUE)]))
        paras.append((None, 4 if k else 0, [(hd, "Georgia", 16, True, WHITE if dark else NAVY)]))
        paras.append((None, 6, [(b, None, 12, False, SUBN if dark else INK)]))
        txt(s, x + 0.28, y + 0.22, w - 0.56, h - 0.3, paras)


# ---------------------------------------------------------------- S1 TITLE
s = new_slide()
slide_bg(s, NAVY)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, 10.23, 0, 3.1, 7.5, NAVY2)
box(s, 10.23, 0, 0.04, 7.5, GOLD)
txt(s, 0.6, 1.02, 9.0, 0.3, [(None, 0, [("ISBMS · PGDM 2025–27 · SEMESTER III · SESSION 10 OF 10", None, 12, True, GOLD)])])
txt(s, 0.6, 1.5, 9.4, 2.4, [
    (None, 0, [("Module 10", "Georgia", 58, True, WHITE)]),
    (None, 10, [("GCC Capstone: One Integrated Financial AI System", "Georgia", 21, False, SUBN)]),
])
box(s, 0.6, 4.28, 4.4, 0.02, GOLD)
txt(s, 0.6, 4.52, 9.0, 1.2, [
    (None, 0, [("Faculty · Agentic AI & Advanced Analytics in Finance", None, 14, True, WHITE)]),
    (None, 4, [("PGDM-SEM3-SPEC-AIFINANCE", None, 12, False, MUTEN)]),
    (None, 3, [("tejasgjadhav.github.io/AIFINANCE", None, 12, True, GOLD)]),
])
txt(s, 10.55, 0.93, 2.4, 0.3, [(None, 0, [("TODAY", None, 10, True, GOLD)])])
txt(s, 10.55, 1.35, 2.5, 1.4, [
    (None, 0, [("Three questions", "Georgia", 15, True, WHITE)]),
    (None, 4, [("How do nine modules", None, 10, False, MUTEN)]),
    (None, 2, [("become one system?", None, 10, False, MUTEN)]),
    (None, 2, [("How is it marked?", None, 10, False, MUTEN)]),
    (None, 2, [("How do I submit it?", None, 10, False, MUTEN)]),
])
txt(s, 10.55, 3.45, 2.5, 1.2, [
    (None, 0, [("Worked example", "Georgia", 15, True, WHITE)]),
    (None, 4, [("A daily market dashboard", None, 10, False, MUTEN)]),
    (None, 2, [("that updates itself every", None, 10, False, MUTEN)]),
    (None, 2, [("weekday morning.", None, 10, False, MUTEN)]),
])
txt(s, 10.55, 4.95, 2.5, 1.6, [
    (None, 0, [("50", "Georgia", 44, True, GOLD)]),
    (None, 4, [("MARKS: REPORT 20,", None, 9, False, MUTEN)]),
    (None, 2, [("VIVA ON YOUR DECK 30", None, 9, False, MUTEN)]),
])
footer(s, 0, dark=True)

# ---------------------------------------------------------------- S2 DELIVERABLES
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 1 · WHAT YOU HAND IN")
title(s, "Two things to submit, one viva to sit")
cards(s, [
    ("DELIVERABLE 1 · 20 MARKS", "The report", "25 to 30 pages, Word or PDF. Times New Roman 12, 1.5 spacing. Similarity 20% or lower. Hard copy too."),
    ("DELIVERABLE 2", "The PowerPoint deck", "8 to 10 slides: objectives, data, analysis and findings. The viva is conducted on this deck."),
    ("THE VIVA · 30 MARKS", "Present and defend", "One internal and one external examiner. A three-minute live demo that shows one failure."),
], cols=3, dark_last=True)
goal_band(s, "The rule behind it all:", "secondary data only. Every number traces to a published source in your bibliography.")
footer(s, 0)

# ---------------------------------------------------------------- S3 FIVE LAYERS
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 1 · CAPSTONE INTEGRATION")
title(s, "One system, built in five layers")
cards(s, [
    ("L1 · INGEST", "Collect the data", "Pull every input from a named public source. Log the source, the URL and the date."),
    ("L2 · EXTRACT", "Find the structure", "Pull tables and numbers out of filings, or index documents for search (RAG)."),
    ("L3 · ENGINE", "Do the maths", "Plain Python, tested against a hand-worked example. This layer earns 20 of the 50 marks."),
    ("L4 · COMMENTARY", "Explain the numbers", "AI writes about the engine's output only, and cites every number it uses."),
    ("L5 · OUTPUT", "Show it", "An app or dashboard the panel can open, plus a document for your annexure."),
    ("THE ONE RULE", "Maths in Python, never in AI", "A language model works on words. A number it made up cannot be defended at the viva."),
], cols=3, dark_last=True)
goal_band(s, "Your nine module outputs are the parts:", "RAG bot, VaR engine, Pillar 3 generator, NAV engine and dashboard become layers of one system.")
footer(s, 0)

# ---------------------------------------------------------------- S4 DASHBOARD AS FIVE LAYERS
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 1 · WORKED EXAMPLE")
title(s, "A daily market dashboard, seen as five layers")
cards(s, [
    ("L1 · INGEST", "Morning inputs", "GIFT Nifty, the previous Nifty close, FII and DII flows, crude oil, USD/INR, India VIX and the morning headlines."),
    ("L2 · EXTRACT", "Numbers from the news", "Pull the figures and the key events out of each headline and report."),
    ("L3 · ENGINE", "Simple calculations", "Expected gap at the open, net institutional flow, and whether volatility is low or high."),
    ("L4 · COMMENTARY", "The pre-market brief", "The Claude API writes a short brief using only those numbers."),
    ("L5 · OUTPUT", "A web page", "A dashboard page, published again every weekday morning."),
    ("WHY IT IS A GOOD MODEL", "It runs without a person", "It updates itself before the market opens. Your capstone should run end to end the same way."),
], cols=3, dark_last=True)
goal_band(s, "Note:", "we use the dashboard to show the design. Your project must not give buy, sell or hold calls without an educational-use disclaimer.")
footer(s, 0)

# ---------------------------------------------------------------- S5 HOW IT RUNS ITSELF
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 1 · HOW IT RUNS ITSELF")
title(s, "Every weekday morning, with no one pressing a button")
flow = [("1", "Schedule", "GitHub Actions starts the job every weekday morning."),
        ("2", "Python script", "Fetches the morning data and runs the calculations."),
        ("3", "Claude API", "Writes the brief from a fixed prompt template."),
        ("4", "New page", "The script writes the updated web page."),
        ("5", "Publish", "The page is saved to GitHub and goes live on GitHub Pages.")]
w, gap = 2.25, 0.22
for i, (n, h, b) in enumerate(flow):
    x = 0.6 + i * (w + gap)
    dark = i == 2
    box(s, x, 1.65, w, 2.3, NAVY if dark else WHITE)
    box(s, x, 1.65, w, 0.05, GOLD)
    txt(s, x + 0.2, 1.82, w - 0.4, 2.05, [
        (None, 0, [(n, "Georgia", 20, True, GOLD)]),
        (None, 2, [(h, "Georgia", 15, True, WHITE if dark else NAVY)]),
        (None, 5, [(b, None, 11, False, SUBN if dark else INK)]),
    ])
    if i < 4:
        txt(s, x + w - 0.02, 2.6, gap + 0.04, 0.4, [(PP_ALIGN.CENTER, 0, [("→", None, 14, True, GOLD)])], align=PP_ALIGN.CENTER)
cards(s, [
    ("GOVERNANCE · MODULE 7", "The key stays secret", "The API key sits in GitHub's secret store, never in the code. The job stops if the key is missing."),
    ("FOR YOUR CAPSTONE", "Show the same discipline", "Keep your work in a public repository, with commits across the whole project period."),
], cols=2, top=4.15, h=2.0, dark_last=True)
goal_band(s, "Cost:", "it uses the per-token API from Module 7. One short brief a day costs very little.")
footer(s, 0)

# ---------------------------------------------------------------- S6 REPORT CHAPTERS
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 1 · THE REPORT")
title(s, "Your system writes most of your report")
cards(s, [
    ("CHAPTERS 6 TO 8", "Summary, introduction, objectives", "Your question, how the work is done by hand today, and your one headline finding."),
    ("CHAPTERS 9 AND 10", "Limits and method", "What you did not cover, your five layers and your data sources. Source: sources.csv."),
    ("CHAPTER 11", "Analysis", "The engine's results and the validation test that proves them right."),
    ("CHAPTERS 12 AND 13", "Findings and recommendation", "What the numbers say, and what a firm should do about it."),
    ("CHAPTER 14", "Bibliography", "Every source you used. A number with no source must be removed."),
    ("CHAPTER 15", "Annexures", "The document your system generated, and your test results."),
], cols=3, dark_last=True)
goal_band(s, "Remember:", "the examiner will pick one number from your report and ask where it came from and how it was calculated.")
footer(s, 0)

# ---------------------------------------------------------------- S7 VIVA DECK / PITCH
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 2 · YOUR VIVA DECK")
title(s, "Build your 8 to 10 slides like a GCC pitch")
cards(s, [
    ("SLIDES 1 AND 2", "The problem", "Your question, and how it is done by hand today."),
    ("SLIDES 3 AND 4", "The solution architecture", "Your data sources and your five layers on one slide."),
    ("SLIDES 5 AND 6", "The engine and its proof", "The calculation, your assumptions and your validation test."),
    ("SLIDE 7", "Live demo", "Three minutes. Include one failure, such as the system refusing a bad input."),
    ("SLIDES 8 AND 9", "Business impact", "Your finding, your recommendation and your limits."),
    ("SLIDE 10", "The team", "Who built what. Both members must be able to explain every layer."),
], cols=3, dark_last=True)
goal_band(s, "Pitch order:", "problem, solution architecture, live demo, business impact, team. This is how GCC hiring panels expect to hear it.")
footer(s, 0)

# ---------------------------------------------------------------- S8 VIVA MARKS
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 3 · HOW THE VIVA IS MARKED")
title(s, "Thirty marks, four things the panel checks")
cards(s, [
    ("5 MARKS", "Is it really yours?", "Where the data came from, why you chose your method, and how well you use your tools."),
    ("10 MARKS", "Do you understand the maths?", "Explain each variable, defend each assumption and walk through each formula."),
    ("5 MARKS", "Do you know the context?", "Link your work to current events, regulation and your sector."),
    ("10 MARKS", "How do you present?", "Stay calm under pressure, speak precisely, and admit what your work cannot do."),
], cols=2, dark_last=True)
goal_band(s, "Mock GCC panel:", "a 15-minute presentation and 10 minutes of questions, scored on depth, communication, portfolio quality and role fit.")
footer(s, 0)

# ---------------------------------------------------------------- S9 HOW TO SUBMIT
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOW TO SUBMIT")
title(s, "Submit on the course website, once per group")
cards(s, [
    ("STEP 1", "Open the page", "Module 10 card, then \"Submit your capstone (report + PPT)\"."),
    ("STEP 2", "Choose your group", "Pick your group number and type either member's roll number."),
    ("STEP 3", "Attach both files", "The report in Word or PDF, and the deck in PowerPoint or PDF. Up to 20 MB each."),
    ("STEP 4", "Submit", "Both files go together. The status table then shows your group as submitted."),
    ("NEED TO CHANGE A FILE?", "Submit again", "Upload both files again. The latest pair is the one marked."),
    ("STILL REQUIRED", "The hard copy", "Hand in the spiral-bound report before the End Semester Examination."),
], cols=3, dark_last=True)
goal_band(s, "Before you submit:", "check your similarity index, your sources.csv and your repository history.")
footer(s, 0)

# ---------------------------------------------------------------- S10 PRACTICE
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "PRACTICE QUESTIONS")
title(s, "Questions the panel is likely to ask")
qs = [
    ("Q1", "What are the five layers of your system?", "Ingest, extract, engine, commentary, output."),
    ("Q2", "Why does the maths live in Python and not in the AI?", "A language model works on words. A number it produced cannot be defended."),
    ("Q3", "Where did this number in your report come from?", "Name the source in sources.csv and walk through the formula."),
    ("Q4", "How do you know your engine is right?", "It reproduces a hand-worked example and a published figure."),
    ("Q5", "Show me your system failing.", "Feed a bad or missing input and show it refuse, with a clear message."),
    ("Q6", "What can your work not do?", "State your limits honestly: data gaps, assumptions, the period covered."),
    ("Q7", "Which regulation applies to your project?", "Name it, for example SEBI rules for investment advice, and how you complied."),
    ("Q8", "How does the dashboard example update itself?", "A scheduled job runs a script, Claude writes the brief, and the page is published."),
]
y = 1.5
for qn_, q, tip in qs:
    box(s, 0.6, y, 12.13, 0.6, WHITE)
    box(s, 0.6, y, 0.07, 0.6, BLUE)
    txt(s, 0.9, y + 0.12, 0.7, 0.4, [(None, 0, [(qn_, "Georgia", 13, True, BLUE)])])
    txt(s, 1.65, y + 0.06, 10.9, 0.55, [
        (None, 0, [(q, None, 11, True, INK)]),
        (None, 1, [("TIP  ", None, 9, True, GOLD), (tip, None, 9.5, False, MUTE)]),
    ])
    y += 0.67
footer(s, 0)

OUT = "slides/Module-10-GCC-Capstone.pptx"
prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
