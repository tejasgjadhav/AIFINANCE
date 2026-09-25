#!/usr/bin/env python3
"""Build Module 6 deck: RegTech & Automated Basel III/IV Compliance.
Five slides: title, Hour 1, Hour 2, Hour 3 lab, practice questions."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

NAVY = RGBColor(0x0F, 0x21, 0x45)
NAVY2 = RGBColor(0x0B, 0x19, 0x36)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLDD = RGBColor(0x8F, 0x71, 0x13)
BLUE = RGBColor(0x3A, 0x5C, 0x9E)
INK = RGBColor(0x1B, 0x2A, 0x41)
BG = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTE = RGBColor(0x5D, 0x6B, 0x7E)
MUTEN = RGBColor(0x8F, 0xA1, 0xC4)  # muted on navy
SUBN = RGBColor(0xC9, 0xD6, 0xEE)   # subtitle on navy
TINT1 = RGBColor(0xE9, 0xEF, 0xF8)
TINT2 = RGBColor(0xD5, 0xE1, 0xF2)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

FOOT_L = "ISBMS · PGDM 2025–27 · AGENTIC AI & ADVANCED ANALYTICS IN FINANCE"


def slide_bg(s, color):
    el = s.background.fill
    el.solid()
    el.fore_color.rgb = color


def box(s, x, y, w, h, fill=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.line.fill.background()
    sh.shadow.inherit = False
    if fill is not None:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    return sh


def txt(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, wrap=True):
    """paras: list of paragraphs; each = (align_or_None, space_before_pt, [runs])
    run = (text, font_name_or_None, size_pt, bold, color)"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in (tf.margin_left, ):
        pass
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for p_align, sp, runs in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = p_align if p_align is not None else align
        if sp:
            p.space_before = Pt(sp)
        for text, fname, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            if fname:
                r.font.name = fname
    return tb


def eyebrow(s, label):
    box(s, 0.6, 0.44, 0.13, 0.13, GOLD)
    txt(s, 0.86, 0.34, 11.83, 0.3, [(None, 0, [(label, None, 11, True, GOLD)])])


def title(s, text):
    txt(s, 0.6, 0.68, 12.13, 0.8, [(None, 0, [(text, "Georgia", 29, True, NAVY)])])


def footer(s, num, dark=False):
    c = MUTEN if dark else MUTE
    txt(s, 0.6, 7.08, 8.0, 0.25, [(None, 0, [(FOOT_L, None, 8, False, c)])])
    txt(s, 10.13, 7.08, 2.6, 0.25, [(PP_ALIGN.RIGHT, 0, [(f"MODULE 6  ·  {num:02d}", None, 8, False, c)])],
        align=PP_ALIGN.RIGHT)


def goal_band(s, lead, rest):
    box(s, 0.6, 6.3, 12.13, 0.66, NAVY)
    box(s, 0.6, 6.3, 0.07, 0.66, GOLD)
    txt(s, 0.9, 6.3, 11.53, 0.66,
        [(None, 0, [(lead + "  ", None, 12, True, GOLD), (rest, None, 12, False, WHITE)])],
        anchor=MSO_ANCHOR.MIDDLE)


def bullet(s, x, y, w, text, size=11):
    txt(s, x, y, w, 0.55, [(None, 0, [("▪  ", None, size, True, GOLD), (text, None, size, False, INK)])])


def new_slide():
    return prs.slides.add_slide(BLANK)


def minibox(s, x, y, label, w=0.62, h=0.34, fill=NAVY, tc=WHITE):
    box(s, x, y, w, h, fill)
    txt(s, x, y, w, h, [(PP_ALIGN.CENTER, 0, [(label, None, 8, True, tc)])],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def arrow(s, x, y, ch="→", w=0.3):
    txt(s, x, y, w, 0.34, [(PP_ALIGN.CENTER, 0, [(ch, None, 13, True, GOLD)])],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def three_cards(s, cards, top=1.72, h=4.2):
    """cards: list of (fill, kick, head, body, example, useline)"""
    xs = [0.6, 4.72, 8.84]
    for x, (fill, kick, head, body, ex, use) in zip(xs, cards):
        dark = fill == NAVY
        box(s, x, top, 3.95, h, fill)
        txt(s, x + 0.26, top + 0.2, 3.4, 0.3, [(None, 0, [(kick, None, 10, True, GOLD if dark else BLUE)])])
        txt(s, x + 0.26, top + 0.52, 3.4, 0.4, [(None, 0, [(head, "Georgia", 15, True, WHITE if dark else NAVY)])])
        txt(s, x + 0.26, top + 1.05, 3.4, 1.55, [(None, 0, [(body, None, 10, False, SUBN if dark else INK)])])
        box(s, x + 0.26, top + 2.62, 3.43, 1.0, WHITE)
        txt(s, x + 0.44, top + 2.72, 3.1, 0.85, [
            (None, 0, [("EXAMPLE", None, 8, True, BLUE)]),
            (None, 2, [(ex, None, 9.5, False, INK)]),
        ])
        txt(s, x + 0.26, top + 3.7, 3.4, 0.45, [(None, 0, [(use, None, 10, True, SUBN if dark else MUTE)])])



# ---------------------------------------------------------------- S1 TITLE
s = new_slide()
slide_bg(s, NAVY)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, 10.23, 0, 3.1, 7.5, NAVY2)
box(s, 10.23, 0, 0.04, 7.5, GOLD)
txt(s, 0.6, 1.02, 9.0, 0.3, [(None, 0, [("ISBMS · PGDM 2025–27 · SEMESTER III · SESSION 6 OF 10", None, 12, True, GOLD)])])
txt(s, 0.6, 1.5, 9.4, 2.4, [
    (None, 0, [("Module 6", "Georgia", 58, True, WHITE)]),
    (None, 10, [("RegTech & Automated Basel III/IV Compliance", "Georgia", 21, False, SUBN)]),
])
box(s, 0.6, 4.28, 4.4, 0.02, GOLD)
txt(s, 0.6, 4.52, 9.0, 1.2, [
    (None, 0, [("Faculty · Agentic AI & Advanced Analytics in Finance", None, 14, True, WHITE)]),
    (None, 4, [("PGDM-SEM3-SPEC-AIFINANCE", None, 12, False, MUTEN)]),
    (None, 3, [("tejasgjadhav.github.io/AIFINANCE", None, 12, True, GOLD)]),
])
txt(s, 10.55, 0.93, 2.4, 0.3, [(None, 0, [("TODAY", None, 10, True, GOLD)])])
txt(s, 10.55, 1.35, 2.5, 1.4, [
    (None, 0, [("Three questions", "Georgia", 15, True, WHITE)]),
    (None, 4, [("Does the bank hold enough capital?", None, 10, False, MUTEN)]),
    (None, 2, [("Can it survive 30 bad days?", None, 10, False, MUTEN)]),
    (None, 2, [("Can a machine write the report?", None, 10, False, MUTEN)]),
])
txt(s, 10.55, 3.25, 2.5, 1.2, [
    (None, 0, [("No heavy maths", "Georgia", 15, True, WHITE)]),
    (None, 4, [("Every ratio today is", None, 10, False, MUTEN)]),
    (None, 2, [("one number divided by", None, 10, False, MUTEN)]),
    (None, 2, [("another number.", None, 10, False, MUTEN)]),
])
txt(s, 10.55, 4.85, 2.5, 1.6, [
    (None, 0, [("1", "Georgia", 44, True, GOLD)]),
    (None, 4, [("PILLAR 3 DOCUMENT, GENERATED", None, 9, False, MUTEN)]),
    (None, 2, [("AND CHECKED, BEFORE YOU LEAVE", None, 9, False, MUTEN)]),
])
footer(s, 1, dark=True)

# ---------------------------------------------------------------- S2 HOUR 1
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 1 · BASEL III IN SIX NUMBERS")
title(s, "Six ratios, six divisions. A regulator reads these first.")
txt(s, 0.6, 1.5, 12.13, 0.3, [(None, 0, [
    ("Real figures: ICICI Bank, Basel III Pillar 3 disclosure, 31 March 2025, consolidated. The minimum is the RBI floor with buffers for a D-SIB.",
     None, 10.5, False, MUTE)])])
tiles = [
    ("CET1 RATIO", "15.81%", "min 8.20%", "Shareholders' equity divided by risk weighted assets. The capital that absorbs a loss first."),
    ("TIER 1 RATIO", "15.81%", "min 9.70%", "CET1 plus additional Tier 1 bonds, over the same RWA. This bank has no AT1 outstanding."),
    ("TOTAL CAPITAL RATIO", "16.41%", "min 11.70%", "Tier 1 plus Tier 2, over RWA. The number the press calls the CRAR."),
    ("LEVERAGE RATIO", "10.37%", "min 4.00%", "Tier 1 capital divided by total exposure with no risk weights. The backstop if the weights are wrong."),
    ("LCR", "≥ 100%", "min 100%", "Liquid assets divided by 30 days of net cash outflows. Can the bank survive a month-long run?"),
    ("NSFR", "≥ 100%", "min 100%", "Stable funding available divided by stable funding required. Is the balance sheet funded for a year?"),
]
for i, (k, v, m, body) in enumerate(tiles):
    x = 0.6 + (i % 3) * 4.12
    y = 1.92 + (i // 3) * 1.62
    w = 3.95 if i % 3 < 2 else 3.89
    box(s, x, y, w, 1.5, WHITE)
    box(s, x, y, w, 0.05, GOLD)
    txt(s, x + 0.26, y + 0.16, 2.2, 0.3, [(None, 0, [(k, None, 9.5, True, BLUE)])])
    txt(s, x + w - 1.6, y + 0.1, 1.35, 0.42, [(PP_ALIGN.RIGHT, 0, [(v, "Georgia", 19, True, NAVY)])], align=PP_ALIGN.RIGHT)
    txt(s, x + w - 1.6, y + 0.5, 1.35, 0.25, [(PP_ALIGN.RIGHT, 0, [(m, None, 8.5, True, GOLDD)])], align=PP_ALIGN.RIGHT)
    txt(s, x + 0.26, y + 0.62, w - 0.5, 0.85, [(None, 0, [(body, None, 9.5, False, INK)])])
box(s, 0.6, 5.22, 12.13, 1.0, NAVY)
box(s, 0.6, 5.22, 0.07, 1.0, GOLD)
txt(s, 0.95, 5.3, 5.6, 0.9, [
    (None, 0, [("WHERE THE MINIMUM COMES FROM", None, 9, True, GOLD)]),
    (None, 3, [("Total capital: 9.0% base + 2.5% conservation buffer + 0.2% D-SIB surcharge = 11.7%. CET1 starts at 5.5%, Tier 1 at 7.0%, and the same buffers sit on top.", None, 10, False, WHITE)]),
])
txt(s, 6.8, 5.3, 5.8, 0.9, [
    (None, 0, [("WHAT AXIOMSL AND ONESUMX DO", None, 9, True, GOLD)]),
    (None, 3, [("They pull the ledger, apply the RBI rule tables, and file the return. AI now watches these six ratios every day instead of once a quarter, and flags the drift before the breach.", None, 10, False, WHITE)]),
])
goal_band(s, "Basel IV / FRTB in one line:",
          "the market-risk number moves from VaR to expected shortfall, and a bank's own model must pass a P&L test desk by desk to keep using it.")
footer(s, 2)

# ---------------------------------------------------------------- S3 HOUR 2
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 2 · AML, KYC AND THE REPORT NOBODY WANTS TO WRITE")
title(s, "Where the machine helps, and where a human still signs")
cards = [
    (WHITE, "TRANSACTION MONITORING", "Rules fire. ML ranks.",
     "A rule fires on every cash deposit above a threshold. Most of those alerts are innocent. A model trained on past cases ranks the alerts, so the analyst opens the worst one first instead of the first one.",
     "10,000 alerts a month. The model puts the 200 that matter at the top of the queue.",
     "The rule still fires. The machine only orders the queue."),
    (WHITE, "KYC DOCUMENT AUTOMATION", "OCR reads. NLP checks.",
     "OCR reads the Aadhaar and PAN. NLP matches the name, date of birth and address against the form. Anything blurred, mismatched or expired goes to a person.",
     "Onboarding drops from three days to twenty minutes for a clean file. The dirty file still takes three days.",
     "The human handles the exceptions, not the volume."),
    (NAVY, "THE SUSPICIOUS TRANSACTION REPORT", "AI drafts. A human signs.",
     "FATF Recommendation 20 requires the report. India files it to FIU-IND under the PMLA. A model can draft the narrative from the alert history. The officer reads it, changes it, and signs it.",
     "One STR drafted in four minutes, then read for twenty. The reading is the job.",
     "No filing leaves without a name on it."),
]
three_cards(s, cards, top=1.55, h=4.2)
box(s, 0.6, 5.82, 12.13, 0.44, TINT1)
box(s, 0.6, 5.82, 0.07, 0.44, BLUE)
txt(s, 0.95, 5.82, 11.5, 0.44, [(None, 0, [
    ("WHAT RBI FINES BANKS FOR UNDER THE KYC MASTER DIRECTION:  ", None, 10, True, BLUE),
    ("accounts opened without full verification, periodic KYC updates missed, and monitoring alerts closed without a review. The fine is the small part. The restriction on onboarding new customers is the expensive part.", None, 10, False, INK)])],
    anchor=MSO_ANCHOR.MIDDLE)
goal_band(s, "The rule for this hour:",
          "automation earns its place by removing the false alerts, never by removing the signature.")
footer(s, 3)

# ---------------------------------------------------------------- S4 THE LAB
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "HOUR 3 · THE LAB")
title(s, "One sentence, one Word document, one check")
box(s, 0.6, 1.58, 12.13, 0.84, NAVY2)
box(s, 0.6, 1.58, 0.07, 0.84, GOLD)
txt(s, 0.95, 1.58, 11.5, 0.84, [
    (None, 0, [("THE SENTENCE WE TYPE", None, 8.5, True, GOLD)]),
    (None, 2, [("“Here is a bank's capital, RWA and leverage data in bank_data.json. Compute CET1, Tier 1, total capital, leverage, LCR and NSFR, "
                "compare each with the RBI minimum, and write a Pillar 3 disclosure in Word with the DF-2 and DF-18 tables.”", "Georgia", 12, True, WHITE)]),
], anchor=MSO_ANCHOR.MIDDLE)
parts = [
    (0.6, 3.95, WHITE, "PART A  ·  30 MIN", "Build it", [
        "Real data: ICICI Bank's published capital and leverage figures, March 2025",
        "Six ratios, each one division",
        "Each ratio next to the RBI floor with buffers",
        "A Word document in the RBI table layout, written by python-docx",
    ]),
    (4.72, 3.95, WHITE, "PART B  ·  15 MIN", "Check it", [
        "The bank printed its own ratios: 15.81, 15.81, 16.41 and 10.37",
        "Your document must land on the same four numbers",
        "A mismatch means a wrong input or a wrong formula",
        "A document that does not reconcile is not issued",
    ]),
    (8.84, 3.89, NAVY, "PART C  ·  15 MIN", "Review it", [
        "Swap documents with the next desk",
        "Paste the two ratio tables into any free AI chat and ask for a commentary",
        "Check every number in its paragraph against your tables",
        "Keep only the sentences that survive",
    ]),
]
for x, w, fill, kick, head, items in parts:
    dark = fill == NAVY
    box(s, x, 2.58, w, 3.34, fill)
    txt(s, x + 0.28, 2.8, w - 0.56, 0.3, [(None, 0, [(kick, None, 9.5, True, GOLD if dark else BLUE)])])
    txt(s, x + 0.28, 3.14, w - 0.56, 0.44, [(None, 0, [(head, "Georgia", 17, True, WHITE if dark else NAVY)])])
    for j2, t in enumerate(items):
        txt(s, x + 0.28, 3.72 + j2 * 0.56, w - 0.56, 0.54,
            [(None, 0, [("▪  ", None, 10, True, GOLD), (t, None, 10, False, SUBN if dark else INK)])])
goal_band(s, "Why Part B matters most:",
          "generating the document takes thirty minutes. Finding out it disagrees with the bank's own number takes five, and it is the only step that makes the document worth anything.")
footer(s, 4)

# ---------------------------------------------------------------- S5 PRACTICE
s = new_slide()
slide_bg(s, BG)
eyebrow(s, "PRACTICE QUESTIONS")
title(s, "Six questions before the exam")
qs = [
    ("Q1", "A bank holds CET1 capital of ₹2,75,000 crore against risk weighted assets of ₹17,40,000 crore. What is its CET1 ratio, and does a D-SIB pass?",
     "15.8%. The D-SIB floor is 5.5% + 2.5% buffer + 0.2% surcharge = 8.2%, so it passes with room to spare."),
    ("Q2", "The minimum total capital ratio for this bank is 11.7%. Show where each part of that number comes from.",
     "9.0% base requirement, plus the 2.5% capital conservation buffer, plus the 0.2% D-SIB surcharge."),
    ("Q3", "Say what question the LCR answers, without using the word ratio.",
     "Can the bank pay 30 days of outflows in a stress from assets it can sell today? At 100% or more, yes."),
    ("Q4", "Basel already has a risk-based capital ratio. Why add a leverage ratio that ignores risk weights?",
     "Risk weights can be wrong or gamed. The leverage ratio is the backstop that does not depend on them."),
    ("Q5", "An AML system fires 10,000 alerts a month and 9,500 are false. What does ML change, and what must it never change?",
     "It ranks the alerts so the analyst reaches the real ones first. It never closes an alert or files a report without a human signature."),
    ("Q6", "Your generated Pillar 3 document shows a total capital ratio of 16.9%. The bank published 16.41%. What do you do?",
     "Do not issue it. Find the wrong input or the wrong formula. Reconciliation to the published number is the gate."),
]
y = 1.58
for qn_, q, tip in qs:
    box(s, 0.6, y, 12.13, 0.74, WHITE)
    box(s, 0.6, y, 0.07, 0.74, BLUE)
    txt(s, 0.9, y + 0.07, 0.7, 0.4, [(None, 0, [(qn_, "Georgia", 14, True, BLUE)])])
    txt(s, 1.7, y + 0.06, 10.8, 0.64, [
        (None, 0, [(q, None, 10.5, True, INK)]),
        (None, 2, [("TIP  ", None, 9, True, GOLD), (tip, None, 9, False, MUTE)]),
    ])
    y += 0.8
box(s, 0.6, 6.42, 12.13, 0.56, NAVY)
box(s, 0.6, 6.42, 0.07, 0.56, GOLD)
txt(s, 0.9, 6.42, 11.5, 0.56, [(None, 0, [
    ("Before Session 7:  ", None, 11, True, GOLD),
    ("open the Pillar 3 disclosure of any listed Indian bank and write down its CET1 ratio and the minimum it quotes against it.", None, 11, False, WHITE)])],
    anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.6, 7.08, 9.0, 0.25, [(None, 0, [("The tip is the shape of the answer, not the whole answer. Two or three sentences each in the exam.", None, 8, False, MUTE)])])
txt(s, 10.13, 7.08, 2.6, 0.25, [(PP_ALIGN.RIGHT, 0, [("MODULE 6  ·  05", None, 8, False, MUTE)])], align=PP_ALIGN.RIGHT)

OUT = "slides/Module-06-RegTech-Basel-Compliance.pptx"
prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
