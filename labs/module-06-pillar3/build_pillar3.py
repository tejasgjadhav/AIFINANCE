"""
Module 6 lab: from one structured data file to the RBI Pillar 3 filing set.

  1. Reads bank_data.json (one block per RBI table, DF-1 to DF-18).
  2. Computes every ratio it can: CET1, Tier 1, total capital, leverage, LCR, NSFR, RWA by risk type.
  3. Checks the computed ratios against the ratios the bank published. A mismatch stops the run.
  4. Writes ONE Word file PER RBI TABLE, in RBI order, into output/. A table the data cannot fill
     is written as a stub that says so. Nothing is invented.
  5. Writes output/00_filing_log.csv, the checklist a compliance officer files against, and one
     combined document for reading.

No API key. Python and python-docx only.   Run:  python3 build_pillar3.py
"""
import json, os, csv, sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = json.load(open("bank_data.json"))
mins, liq = d["rbi_minimums_pct"], d["liquidity_ILLUSTRATIVE"]
OUT = "output"; os.makedirs(OUT, exist_ok=True)
def mn(x): return f"{x:,.1f}"
def pc(x): return f"{x:.2f}%"

# ---------- 1. compute ----------
c2 = d["DF-2"]; cap, req, R = c2["capital"], c2["capital_requirement"], c2["total_rwa"]
tier1 = cap["cet1_capital"] + cap["additional_tier1_capital"]
total_cap = tier1 + cap["tier2_capital"]
buf = mins["capital_conservation_buffer"] + (mins["dsib_surcharge"] if d["is_dsib"] else 0)
minimum = {"CET1 ratio": mins["cet1"] + buf, "Tier 1 ratio": mins["tier1"] + buf, "Total capital ratio": mins["total_capital"] + buf,
           "Leverage ratio": mins["leverage_dsib"] if d["is_dsib"] else mins["leverage_other"], "LCR": mins["lcr"], "NSFR": mins["nsfr"]}
l18 = d["DF-18"]
ratios = {"CET1 ratio": cap["cet1_capital"] / R * 100, "Tier 1 ratio": tier1 / R * 100, "Total capital ratio": total_cap / R * 100,
          "Leverage ratio": l18["tier1_capital"] / l18["total_exposure"] * 100,
          "LCR": liq["hqla"] / liq["net_cash_outflows_30d"] * 100,
          "NSFR": liq["available_stable_funding"] / liq["required_stable_funding"] * 100}
rwa_by = {k: req[k] / (minimum["Total capital ratio"] / 100) for k in ("credit_risk_total", "market_risk", "operational_risk")}

# ---------- 2. reconcile before writing anything ----------
pub = c2["published_ratios_pct"]
checks = [("CET1 ratio", pub["cet1_crar"]), ("Tier 1 ratio", pub["tier1_crar"]), ("Total capital ratio", pub["total_crar"]),
          ("Leverage ratio", l18["published_leverage_ratio_pct"])]
print("RECONCILIATION against the bank's published ratios")
bad = 0
for name, p in checks:
    ok = abs(round(ratios[name], 2) - p) <= 0.01; bad += (not ok)
    print(f"  {name:22s} computed {ratios[name]:6.2f}%  published {p:6.2f}%  RBI min {minimum[name]:5.2f}%  {'MATCH' if ok else 'MISMATCH'}")
d17 = d["DF-17"]
s17 = sum(d17[k] for k in ("total_consolidated_assets", "adj_investments_in_banking_financial_insurance", "adj_fiduciary_assets", "adj_derivatives", "adj_sfts", "adj_off_balance_sheet", "other_adjustments"))
ok17 = abs(s17 - d17["leverage_ratio_exposure"]) < 0.2; bad += (not ok17)
print(f"  {'DF-17 adds up':22s} {mn(s17)} vs {mn(d17['leverage_ratio_exposure'])}  {'MATCH' if ok17 else 'MISMATCH'}")
d4 = d["DF-4"]["exposure_by_risk_weight_after_crm"]
ok4 = abs(d4["below_100pct"] + d4["at_100pct"] + d4["above_100pct"] - d4["total"]) < 0.2; bad += (not ok4)
print(f"  {'DF-4 adds up':22s} {'MATCH' if ok4 else 'MISMATCH'}")
if bad:
    sys.exit("STOP: a computed figure does not reconcile to the published one. Fix the input or the formula. Nothing was filed.")

# ---------- 3. document helpers ----------
def new_doc(code, title, status, pages):
    doc = Document(); doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(f"Table {code}: {title}", 0)
    p = doc.add_paragraph(); r = p.add_run(f"{d['bank']}  |  {d['as_of']}  |  {d['unit']}"); r.bold = True
    p = doc.add_paragraph(); r = p.add_run(f"Source: {d['source']}, pages {pages}. Data status: {status}."); r.italic = True; r.font.size = Pt(9)
    return doc
def para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; return p
def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
    for i, x in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(x).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = x
            if i > 0: cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()
def stub(doc, code):
    para(doc, f"This table is not populated. bank_data.json carries no figures for {code}. The disclosure text and figures "
              f"must be taken from the bank's PDF, pages {d[code]['pages']}, before this file can be filed.", italic=True)

# ---------- 4. one file per RBI table, in order ----------
def build(code):
    b = d[code]; doc = new_doc(code, b["title"], b["status"], b["pages"])
    if code == "DF-1":
        para(doc, "Group entities named in the disclosure:", bold=True)
        table(doc, ["Entity", "Consolidation treatment"], [[e, "see PDF pages 1-8"] for e in b["group_entities_named"]])
        para(doc, b["_note"], italic=True)
    elif code == "DF-2":
        para(doc, f"The Bank is a Domestic Systemically Important Bank. Minimum total capital ratio including the {mins['capital_conservation_buffer']}% "
                  f"capital conservation buffer and the {mins['dsib_surcharge']}% D-SIB surcharge is {pc(minimum['Total capital ratio'])} of risk weighted assets.")
        para(doc, "Capital requirements", bold=True)
        table(doc, ["Item", "Capital required", "Risk weighted assets"], [
            ["(b) Credit risk: standardised approach", mn(req["credit_risk_standardised"]), ""],
            ["      of which securitisation exposures", mn(req["securitisation"]), ""],
            ["      Credit risk total", mn(req["credit_risk_total"]), mn(rwa_by["credit_risk_total"])],
            ["(c) Market risk: standardised duration approach", mn(req["market_risk"]), mn(rwa_by["market_risk"])],
            ["(d) Operational risk: basic indicator approach", mn(req["operational_risk"]), mn(rwa_by["operational_risk"])],
            ["Total capital requirement (b + c + d)", mn(req["total"]), mn(R)]])
        para(doc, "Capital ratios", bold=True)
        table(doc, ["Ratio", "Computed", "Published", "RBI minimum incl. buffers", "Headroom"], [
            [k, pc(ratios[k]), pc(p), pc(minimum[k]), f"{ratios[k]-minimum[k]:.2f} pp"] for k, p in checks[:3]])
        para(doc, "Composition of capital funds", bold=True)
        table(doc, ["Component", "Amount"], [["Common Equity Tier 1", mn(cap["cet1_capital"])], ["Additional Tier 1", mn(cap["additional_tier1_capital"])],
              ["Tier 1", mn(tier1)], ["Tier 2", mn(cap["tier2_capital"])], ["Total capital funds", mn(total_cap)], ["Total risk weighted assets", mn(R)]])
        para(doc, "Subsidiary ratios as published: ICICI Bank UK PLC CET1 %s / total %s; ICICI Bank Canada CET1 %s / total %s; standalone bank CET1 %s / total %s."
             % tuple(pc(pub[k]) for k in ("uk_cet1", "uk_total", "canada_cet1", "canada_total", "standalone_cet1", "standalone_total")))
    elif code == "DF-3":
        n, m, pr, npi = b["npa"], b["npa_movement"], b["provisions"], b["npi"]
        para(doc, "Non-performing assets", bold=True)
        table(doc, ["Item", "Amount"], [["Sub-standard", mn(n["substandard"])], ["Doubtful", mn(n["doubtful"])], ["Loss", mn(n["loss"])],
              ["Gross NPAs", mn(n["gross_npa"])], ["Net NPAs", mn(n["net_npa"])], ["Gross NPA ratio", pc(n["gross_npa_ratio_pct"])], ["Net NPA ratio", pc(n["net_npa_ratio_pct"])]])
        para(doc, "Movement of gross NPAs", bold=True)
        table(doc, ["Item", "Amount"], [["Opening balance", mn(m["opening"])], ["Additions", mn(m["additions"])], ["Reductions", mn(m["reductions"])], ["Closing balance", mn(m["closing"])]])
        para(doc, "Provisions and non-performing investments", bold=True)
        table(doc, ["Item", "Amount"], [["Specific provisions, opening", mn(pr["specific_opening"])], ["Specific provisions, closing", mn(pr["specific_closing"])],
              ["General provisions, closing", mn(pr["general_closing"])], ["Gross NPIs", mn(npi["gross"])], ["Provisions for NPIs", mn(npi["provisions"])], ["Net NPIs", mn(npi["net"])]])
    elif code == "DF-4":
        e = b["exposure_by_risk_weight_after_crm"]
        para(doc, "Credit exposures by risk weight, after credit risk mitigation", bold=True)
        table(doc, ["Risk weight bucket", "Exposure"], [["Less than 100%", mn(e["below_100pct"])], ["100%", mn(e["at_100pct"])], ["More than 100%", mn(e["above_100pct"])], ["Total", mn(e["total"])]])
    elif code == "DF-6":
        para(doc, "Capital required for securitisation exposures under the standardised approach", bold=True)
        table(doc, ["Item", "Amount"], [["Capital required, securitisation", mn(b["capital_required_securitisation"])]])
        para(doc, "The exposure-level breakdown on pages 32-38 is not in bank_data.json.", italic=True)
    elif code == "DF-7":
        c = b["capital_required"]; para(doc, f"Approach: {b['approach']}.")
        table(doc, ["Risk", "Capital required"], [["Interest rate risk", mn(c["interest_rate_risk"])], ["Foreign exchange risk, including gold", mn(c["foreign_exchange_incl_gold"])],
              ["Equity position risk", mn(c["equity_position_risk"])], ["Total", mn(c["total"])]])
    elif code == "DF-8":
        para(doc, f"Approach: {b['approach']}."); table(doc, ["Item", "Amount"], [["Capital required for operational risk", mn(b["capital_required"])]])
    elif code == "DF-9":
        para(doc, f"Impact of a {b['shock_bps']} bps downward parallel shift in interest rates. {b['_note']}")
        table(doc, ["Currency", "Earnings (NII) impact", "Economic value (EVE) impact"],
              [[k, mn(b["nii_impact"][k]), mn(b["eve_impact"][k])] for k in ("INR", "USD", "Others", "Total")])
    elif code == "DF-16":
        table(doc, ["Item", "Amount"], [["Book value of banking-book equity investments", mn(b["book_value"])], ["Market value", mn(b["market_value"])],
              ["Realised gain from sale in the period", mn(b["realised_gain_in_period"])], ["Unrealised gain or loss", mn(b["unrealised_gain_loss"])], ["Latent revaluation gain", mn(b["latent_revaluation_gain"])]])
    elif code == "DF-17":
        table(doc, ["Item", "Amount"], [["Total consolidated assets as per published financial statements", mn(b["total_consolidated_assets"])],
              ["Adjustment for investments in banking, financial, insurance or commercial entities", mn(b["adj_investments_in_banking_financial_insurance"])],
              ["Adjustment for fiduciary assets", mn(b["adj_fiduciary_assets"])], ["Adjustments for derivative financial instruments", mn(b["adj_derivatives"])],
              ["Adjustment for securities financing transactions", mn(b["adj_sfts"])], ["Adjustment for off-balance sheet items", mn(b["adj_off_balance_sheet"])],
              ["Other adjustments", mn(b["other_adjustments"])], ["Leverage ratio exposure", mn(b["leverage_ratio_exposure"])]])
        para(doc, "Check: the rows add to the leverage ratio exposure. " + ("Passed." if ok17 else "FAILED."), italic=True)
    elif code == "DF-18":
        table(doc, ["Item", "Amount"], [["On-balance sheet items, excluding derivatives and SFTs", mn(b["on_balance_sheet_items_excl_derivatives_sfts"])],
              ["Asset amounts deducted in determining Tier 1 capital", mn(b["asset_amounts_deducted"])], ["Total on-balance sheet exposures", mn(b["on_balance_sheet_exposure"])],
              ["Total derivative exposures", mn(b["derivative_exposure"])], ["Total securities financing transaction exposures", mn(b["sft_exposure"])],
              ["Off-balance sheet items", mn(b["off_balance_sheet_items"])], ["Total exposures", mn(b["total_exposure"])], ["Tier 1 capital", mn(b["tier1_capital"])],
              ["Basel III leverage ratio, computed", pc(ratios["Leverage ratio"])], ["Basel III leverage ratio, published", pc(b["published_leverage_ratio_pct"])], ["RBI minimum, D-SIB", pc(minimum["Leverage ratio"])]])
    else:
        stub(doc, code)
    return doc

codes = [f"DF-{i}" for i in range(1, 19)]
log = []
print("\nFILING, in RBI order")
for code in codes:
    b = d[code]; doc = build(code)
    fn = f"{OUT}/{code.replace('-', '-') if False else code}_{b['title'].replace(':', '').replace(' ', '_')}.docx"
    fn = f"{OUT}/{int(code[3:]):02d}_{code}_{b['title'].replace(':', '').replace(' ', '_')}.docx"
    doc.save(fn)
    log.append([code, b["title"], b["status"], b["pages"], os.path.basename(fn)])
    print(f"  {code:6s} {b['status']:12s} -> {os.path.basename(fn)}")

# liquidity appendix (not an RBI DF table)
doc = new_doc("Appendix", liq["title"], "ILLUSTRATIVE inputs", "n/a")
para(doc, liq["_note"], italic=True)
table(doc, ["Item", "Amount"], [["High quality liquid assets", mn(liq["hqla"])], ["Net cash outflows over 30 days", mn(liq["net_cash_outflows_30d"])],
      ["Liquidity Coverage Ratio", pc(ratios["LCR"])], ["Available stable funding", mn(liq["available_stable_funding"])], ["Required stable funding", mn(liq["required_stable_funding"])],
      ["Net Stable Funding Ratio", pc(ratios["NSFR"])], ["RBI minimum for each", pc(100.0)]])
doc.save(f"{OUT}/99_Appendix_LCR_NSFR_illustrative.docx"); log.append(["Appendix", liq["title"], "illustrative", "n/a", "99_Appendix_LCR_NSFR_illustrative.docx"])

with open(f"{OUT}/00_filing_log.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["table", "title", "data_status", "pdf_pages", "file"]); w.writerows(log)

# combined reading copy
combo = Document()
combo.add_heading("Basel III Pillar 3 Disclosures: filing set", 0)
para(combo, f"{d['bank']}  |  {d['as_of']}  |  {d['unit']}", bold=True)
para(combo, "One section per RBI table, generated from bank_data.json. See 00_filing_log.csv for the status of each.", italic=True)
for code in codes:
    src = Document(f"{OUT}/{int(code[3:]):02d}_{code}_{d[code]['title'].replace(':', '').replace(' ', '_')}.docx")
    for el in src.element.body:
        if el.tag.endswith("sectPr"): continue
        combo.element.body.append(el)
combo.save("Pillar3_Disclosure_ICICI_Mar2025.docx")
n_ok = sum(1 for r in log if r[2] == "complete"); n_part = sum(1 for r in log if r[2] == "partial"); n_stub = sum(1 for r in log if r[2] == "not_in_data")
print(f"\n{n_ok} tables complete, {n_part} partial, {n_stub} stubs, out of 18. Log: {OUT}/00_filing_log.csv. Combined copy: Pillar3_Disclosure_ICICI_Mar2025.docx")
