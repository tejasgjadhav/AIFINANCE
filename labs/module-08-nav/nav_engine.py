"""
Module 8 lab: a daily NAV engine with a reconciliation check.

Run:  python3 nav_engine.py

Built-in Python only. No internet, no API key, no pip install.
The fund is made up for teaching. Holdings are named Stock A to Stock J and the
prices are illustrative, not market data.

What it does, in order:
  1. Prices the fund from its own books and works out NAV per unit.
  2. Reconciles the books against the custodian (units and cash) and the
     transfer agent (units outstanding). Every difference is a "break".
  3. Says what each break most likely is, and what entry would fix it.
  4. Works out the corrected NAV once the breaks are fixed.
  5. Writes nav_report.csv and breaks.csv, and prints a prompt you can paste
     into any free AI chat to draft the fund manager's commentary.
  6. Optional: if the openpyxl library is available (it is built into Google
     Colab), it also writes nav_report.xlsx, a formatted report laid out like a
     fund's NAV and portfolio disclosure. Without openpyxl this step is skipped.
"""
import csv, os

HERE = os.path.dirname(os.path.abspath(__file__))


def read_csv(name):
    with open(os.path.join(HERE, name), newline="") as f:
        return list(csv.DictReader(f))


def rupees(x):
    return f"Rs {x:,.2f}"


def nav4(x):
    return f"Rs {x:,.4f}"


books = read_csv("fund_books.csv")
custodian = {r["holding"]: float(r["units"]) for r in read_csv("custodian.csv")}
info = {r["item"]: float(r["value"]) for r in read_csv("fund_info.csv")}

# ---- 1. NAV from the fund's own books
market_value = sum(float(r["units"]) * float(r["price"]) for r in books)
cash = info["cash_per_books"]
expenses = info["accrued_expenses"]
units_out = info["units_outstanding_per_books"]
net_assets = market_value + cash - expenses
nav = net_assets / units_out

print("STEP 1  NAV FROM THE FUND'S BOOKS")
print(f"  Market value of 10 holdings   {rupees(market_value):>22}")
print(f"  Add cash                      {rupees(cash):>22}")
print(f"  Less accrued expenses         {rupees(-expenses):>22}")
print(f"  Net assets                    {rupees(net_assets):>22}")
print(f"  Units outstanding             {units_out:>22,.0f}")
print(f"  NAV per unit                  {nav4(nav):>22}")

# ---- 2. Reconciliation: books vs custodian vs transfer agent
breaks = []
for r in books:
    h, u_books = r["holding"], float(r["units"])
    u_cust = custodian.get(h, 0.0)
    if u_cust != u_books:
        ratio = u_cust / u_books if u_books else 0
        if ratio and float(ratio).is_integer():
            cause = f"Units are exactly {ratio:.0f}x at the custodian: a stock split not booked in the fund's books"
            fix = f"Book the split: units {u_books:,.0f} -> {u_cust:,.0f}, price divided by {ratio:.0f}. Market value does not change."
        else:
            cause = "Units differ: a trade booked on one side only, or a settlement failure"
            fix = "Check the trade blotter and the settlement report before changing anything"
        breaks.append(["Units held", h, u_books, u_cust, u_cust - u_books, cause, fix])

cash_cust = custodian.get("CASH", 0.0)
if round(cash_cust - cash, 2) != 0:
    breaks.append(["Cash", "Custodian account", cash, cash_cust, cash_cust - cash,
                   "Custodian holds more cash than the books: most often a dividend received but not booked",
                   f"Check the dividend notice, then book income of {rupees(cash_cust - cash)}"])

ta_units = info["units_outstanding_per_transfer_agent"]
if ta_units != units_out:
    breaks.append(["Units outstanding", "Transfer agent register", units_out, ta_units, ta_units - units_out,
                   "The transfer agent shows fewer units: a redemption processed by the TA but not in the fund's books",
                   f"Book the redemption of {units_out - ta_units:,.0f} units and the cash payable to the investor"])

print("\nSTEP 2  RECONCILIATION: BOOKS vs CUSTODIAN vs TRANSFER AGENT")
if not breaks:
    print("  No breaks. Books, custodian and transfer agent agree.")
for i, b in enumerate(breaks, 1):
    print(f"  Break {i}: {b[0]} - {b[1]}")
    print(f"     books {b[2]:,.2f}   other side {b[3]:,.2f}   difference {b[4]:,.2f}")
    print(f"     likely cause: {b[5]}")
    print(f"     suggested fix: {b[6]}")

# ---- 3. Corrected NAV
# The split leaves market value unchanged. The dividend adds cash. The redemption
# removes units and creates a payable equal to units x NAV.
cash_fixed = cash_cust
redeemed = units_out - ta_units
nav_pre = (market_value + cash_fixed - expenses) / units_out
payable = redeemed * nav_pre
net_fixed = market_value + cash_fixed - expenses - payable
nav_fixed = net_fixed / ta_units

print("\nSTEP 3  NAV AFTER THE BREAKS ARE FIXED")
print(f"  Net assets                    {rupees(net_fixed):>22}")
print(f"  Units outstanding             {ta_units:>22,.0f}")
print(f"  NAV per unit                  {nav4(nav_fixed):>22}")
print(f"  Change from the unreconciled NAV  {nav4(nav_fixed - nav)} per unit")

# ---- 4. Files
with open(os.path.join(HERE, "nav_report.csv"), "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["line", "before reconciliation", "after reconciliation"])
    w.writerow(["Market value", round(market_value, 2), round(market_value, 2)])
    w.writerow(["Cash", round(cash, 2), round(cash_fixed, 2)])
    w.writerow(["Accrued expenses", round(-expenses, 2), round(-expenses, 2)])
    w.writerow(["Redemption payable", 0, round(-payable, 2)])
    w.writerow(["Net assets", round(net_assets, 2), round(net_fixed, 2)])
    w.writerow(["Units outstanding", units_out, ta_units])
    w.writerow(["NAV per unit", round(nav, 4), round(nav_fixed, 4)])
with open(os.path.join(HERE, "breaks.csv"), "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["type", "item", "books", "other side", "difference", "likely cause", "suggested fix"])
    w.writerows(breaks)
print("\nWrote nav_report.csv and breaks.csv. Open them in Excel or Google Sheets.")

# ---- 5. Commentary prompt for any free AI chat
top = sorted(books, key=lambda r: float(r["units"]) * float(r["price"]), reverse=True)[:3]
print("\nSTEP 4  PASTE THIS INTO ANY FREE AI CHAT FOR THE COMMENTARY")
print("-" * 70)
print("You are a fund accountant. Write a five-sentence daily NAV commentary for")
print("an Indian equity mutual fund. Use only the numbers below. Do not add any number")
print("that is not listed.")
print(f"NAV per unit before reconciliation: {nav4(nav)}")
print(f"NAV per unit after reconciliation: {nav4(nav_fixed)}")
print(f"Net assets after reconciliation: {rupees(net_fixed)}")
print(f"Breaks found: {len(breaks)} (a stock split, a dividend not booked, a redemption not booked)")
print("Largest holdings: " + ", ".join(f"{r['holding']} ({r['sector']})" for r in top))
print("-" * 70)
print("Then check every number in the AI's answer against nav_report.csv.")


# ---- 6. Optional formatted Excel report (needs openpyxl, which Google Colab already has)
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("\nopenpyxl is not installed, so nav_report.xlsx was skipped. Run this in Google Colab to get it.")
else:
    NAVY_FILL = PatternFill("solid", fgColor="0F2145")
    WHITE_B = Font(bold=True, color="FFFFFF")
    BOLD = Font(bold=True)
    thin = Side(style="thin", color="D5E1F2")
    RS = '"Rs "#,##,##0.00'
    wb = Workbook()
    ws = wb.active
    ws.title = "NAV report"
    ws["A1"] = "Illustrative Equity Fund  |  Daily NAV and portfolio statement, after reconciliation"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "A made-up fund for teaching. Holdings are Stock A to Stock J and prices are illustrative."
    ws["A2"].font = Font(italic=True, color="5D6B7E")

    ws.append([])
    ws.append(["PORTFOLIO"])
    ws.cell(ws.max_row, 1).font = BOLD
    hdr = ["Holding", "Sector", "Units", "Price", "Market value", "% of net assets"]
    ws.append(hdr)
    for c in range(1, 7):
        ws.cell(ws.max_row, c).font = WHITE_B
        ws.cell(ws.max_row, c).fill = NAVY_FILL
    # Show the portfolio after reconciliation: a booked split means the custodian's units at the split-adjusted price.
    fixed_rows = []
    for r in books:
        u, pr = float(r["units"]), float(r["price"])
        u_c = custodian.get(r["holding"], u)
        if u_c != u and u and (u_c / u).is_integer():
            pr, u = pr * u / u_c, u_c
        fixed_rows.append((r["holding"], r["sector"], u, pr))
    for h, sec, u, pr in sorted(fixed_rows, key=lambda x: x[2] * x[3], reverse=True):
        mv = u * pr
        ws.append([h, sec, u, pr, mv, mv / net_fixed])
        row = ws.max_row
        ws.cell(row, 3).number_format = "#,##,##0"
        ws.cell(row, 4).number_format = RS
        ws.cell(row, 5).number_format = RS
        ws.cell(row, 6).number_format = "0.00%"

    ws.append([])
    ws.append(["NAV CALCULATION", "", "", "", "Before reconciliation", "After reconciliation"])
    for c in range(1, 7):
        ws.cell(ws.max_row, c).font = WHITE_B
        ws.cell(ws.max_row, c).fill = NAVY_FILL
    lines = [
        ("Holdings at market price", market_value, market_value),
        ("Cash", cash, cash_fixed),
        ("Less accrued expenses", -expenses, -expenses),
        ("Less redemption payable", 0.0, -round(payable, 2)),
        ("Net assets", net_assets, net_fixed),
        ("Units outstanding", units_out, ta_units),
        ("NAV per unit", nav, nav_fixed),
    ]
    for label, before, after in lines:
        ws.append([label, "", "", "", before, after])
        row = ws.max_row
        fmt = "#,##,##0" if label == "Units outstanding" else ('"Rs "0.0000' if label == "NAV per unit" else RS)
        for c in (5, 6):
            ws.cell(row, c).number_format = fmt
        if label in ("Net assets", "NAV per unit"):
            for c in (1, 5, 6):
                ws.cell(row, c).font = BOLD

    ws.append([])
    ws.append(["Prepared by the fund accountant. Reviewed and signed by: ____________________"])

    ws2 = wb.create_sheet("Breaks")
    ws2.append(["Type", "Item", "Books", "Other side", "Difference", "Likely cause", "Suggested fix"])
    for c in range(1, 8):
        ws2.cell(1, c).font = WHITE_B
        ws2.cell(1, c).fill = NAVY_FILL
    for b in breaks:
        ws2.append(b)
        for c in (3, 4, 5):
            ws2.cell(ws2.max_row, c).number_format = "#,##,##0.00"

    for sheet, widths in ((ws, [28, 16, 14, 16, 22, 22]), (ws2, [18, 24, 16, 16, 16, 60, 60])):
        for i, w in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + i)].width = w
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 6 and sheet is ws2)
    wb.save(os.path.join(HERE, "nav_report.xlsx"))
    print("\nWrote nav_report.xlsx, a formatted NAV and portfolio report with a Breaks sheet.")
